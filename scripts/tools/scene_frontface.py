# -*- coding: utf-8 -*-
"""场景资产图「正视」评判与改写工具（2026-10-13 用户口径）。

【流程定位】场景提示词**正常生成**（不带正视强制规则），生成后用本工具做两件事：
  audit   —— 逐场景评判「需要正视吗」：功能主体明确且立面型对象（床/门/柜/书架/货架/
             电视墙/灶台/洗手台+镜…）→ 建议正视；空间型/纵深型/无明确主体
             （街道/广场/大厅/废墟/墓园/走廊/外景全貌）→ 不需正视。
             产出清单分两组，**不需正视的显式列出**——用户后续想改随时 rewrite 单个场景。
  rewrite —— 把指定场景的三版逐版改写成正视版（功能主体正面朝镜头、水平平视、一点透视，
             其他要素原样保留），输出新 docx，不覆盖原文件。

用法：
  python tools/scene_frontface.py audit  --docx <场景提示词.docx> [--out <审计.md>]
  python tools/scene_frontface.py rewrite --docx <场景提示词.docx> --scenes 场景A,场景B [--out <新docx>]

走 ds-free-api 反代（config.py 现配置），调用免费但受风控：串行 + 失败重试由 ds_api 统一处理。
"""
import argparse
import json
import re
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "src"))

from docx import Document  # noqa: E402
from docx_utils import save_text_docx  # noqa: E402
from ds_api import call_ds_api  # noqa: E402

HEAD_RE = re.compile(r"^(?P<name>.+?)\s*-\s*方案(?P<n>[一二三1234]+)（(?P<codes>[^）]*)）")

AUDIT_SYSTEM = """你是场景资产图的构图评审。给定一个场景的三份提示词，判断这个场景出图时
是否应该用「主体正对镜头」的正视机位。

判断规则：
1. 先找出场景的「功能主体」——人物在这个空间里主要使用的固定对象
   （卧室=床、书房=书桌、客厅=沙发与电视墙、厨房=灶台、玄关/门厅=门、
   仓库=货架阵列、盥洗室=洗手台与镜子、店堂=柜台，以此类推）。
2. 主体是「立面型」对象（有明确正面：床头板、门板、柜门、屏幕、台面）→ frontface = true，
   正视有资产价值。
3. 空间型 / 环型 / 纵深型场景（街道、广场、大厅、废墟、墓园、走廊、外景建筑全貌、
   天台等），或找不出明确功能主体 → frontface = false，正视反而让资产显得奇怪。
只输出一行 JSON，不要多余文字：
{"subject": "功能主体名称，没有则填 无", "frontface": true, "reason": "一句话理由"}"""

REWRITE_SYSTEM = """你是场景提示词改写器。把给定的一份场景提示词改写成「正视图」构图，
改写后直接输出提示词全文。不要标注方案号，不要输出任何别的文字。

【正视图的定义（最高准则）】画面是「横平竖直」的构图：空间的主要结构线（地平线、墙线、
路缘线、家具长边、床沿、柜线）要么与画幅边线平行/垂直，要么沿纵深方向对称收敛于画幅中心；
**绝不出现斜向穿画面的构图**（没有斜角、没有偏轴、没有主体斜置）。

改写规则：
1. 先确定场景的功能主体（人物主要使用的固定对象，如床/门/柜/书架/货架/柜台/电视墙/灶台等）。
2. 在提示词开头注入一句硬构图约束（必须含以下全部要点，用摄影术语写死）：
   「正视图构图——相机正对【主体】正面，主体居于画面中央、左右基本对称；
   主体的正立面与画面平行（与画幅边线正交）、不产生侧向透视；一点透视；机位水平、不俯视不仰视；
   画面横平竖直，严禁斜角构图，严禁主体侧置或偏向一侧，严禁从房间角落斜看。」
3. 【判别顺序——先家具后纵深】先找**家具型功能主体**（床/书桌/柜/书架/货架/柜台/门/电视墙）：
   只要有，就一律用它（按第 2 条：主体居中、正立面平行画面）。**房间即使又长又窄，也不得
   因为纵深就把主体丢到画面一侧**。卧室/客房的主体永远是**床**：机位从床尾方向正对床，
   床头板所在的墙与画面平行（正立面），床身沿纵深朝镜头伸出、居中对称，两侧床头柜对称可见。
   **仅当场景是通道型空间**（马路、街道、林荫路、走廊、隧道、车厢通道）**且没有家具型主体**时，
   才用中轴规则：「正视图构图——相机立于道路/空间中轴线上正向拍摄，路面/空间中轴笔直延伸向
   画面深处、对准画幅中心，两侧景物左右对称；一点透视；机位水平、不俯视不仰视；
   严禁把道路拍成斜穿画面，严禁偏轴斜拍。」
4. 其余一切内容严格保持原样：空间结构、家具摆位、材质、颜色、破坏痕迹、文字标识、
   画风与画质后缀、空镜声明，一律不动，只调整机位与构图描述。
5. 如果这份提示词本来就接近正视构图，也按第 2/3 条把构图句补全后返回。"""


def read_scene_docx(path: Path):
    """读场景提示词 docx → {场景名: {"codes":…, "variants": {方案号: 提示词}}}"""
    doc = Document(str(path))
    out = {}
    for p in doc.paragraphs:
        style = (p.style.name or "").lower()
        if not style.startswith("heading"):
            continue
        m = HEAD_RE.match(p.text.strip())
        if not m:
            continue
        name = m.group("name").strip()
        n = m.group("n")
        rec = out.setdefault(name, {"codes": m.group("codes"), "variants": {}})
        rec["variants"][n] = None        # 占位，正文在后面的段落
    # 第二遍：把每条 heading 后面的第一个非空段落作为该方案正文
    cur = None
    for p in doc.paragraphs:
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            m = HEAD_RE.match(p.text.strip())
            cur = m.group("name").strip() if m else None
            continue
        if cur and p.text.strip():
            n = None
            # 找该场景还没填到正文的方案号
            for k, v in out[cur]["variants"].items():
                if v is None:
                    n = k
                    break
            if n:
                out[cur]["variants"][n] = p.text.strip()
            cur = None
    return out


def audit(docx_path: Path, out_path: Path):
    scenes = read_scene_docx(docx_path)
    if not scenes:
        print(f"!! {docx_path.name} 里没解析出任何「场景 - 方案N（…）」结构")
        return 2
    print(f"共 {len(scenes)} 个场景，逐个评判（走反代，串行）…\n")
    rows = []
    for i, (name, rec) in enumerate(scenes.items(), 1):
        joined = "\n\n".join(f"方案{k}：{v}" for k, v in rec["variants"].items() if v)
        try:
            text = call_ds_api(AUDIT_SYSTEM, f"场景名：{name}\n\n{joined}",
                               temperature=0.0, max_tokens=300)
        except Exception as e:
            print(f"[{i}/{len(scenes)}] {name} —— 评判失败: {e!r}")
            rows.append({"name": name, "subject": "?", "frontface": None,
                         "reason": f"评判失败：{e!r}"})
            continue
        # 剥 JSON 壳
        m = re.search(r"\{.*\}", text, re.S)
        try:
            j = json.loads(m.group(0)) if m else json.loads(text)
        except Exception:
            j = {"subject": "?", "frontface": None, "reason": f"JSON 解析失败: {text[:120]}"}
        j.setdefault("name", name)
        rows.append({"name": name, **{k: j.get(k) for k in ("subject", "frontface", "reason")}})
        tag = "正视" if j.get("frontface") is True else ("不需正视" if j.get("frontface") is False else "存疑")
        print(f"[{i}/{len(scenes)}] {name} —— {tag}（主体: {j.get('subject')}）")

    yes = [r for r in rows if r["frontface"] is True]
    no = [r for r in rows if r["frontface"] is False]
    bad = [r for r in rows if r["frontface"] is None]

    lines = [f"# 场景正视审计（{docx_path.name}）", "",
             f"来源：`{docx_path}`；评判口径：功能主体明确的立面型场景建议正视，"
             "空间型/纵深型/无主体场景不需正视。", ""]
    lines.append(f"## ✅ 建议正视（{len(yes)} 个）—— 可用 rewrite 改写")
    for r in yes:
        lines.append(f"- **{r['name']}** —— 主体：{r['subject']}；{r['reason']}")
    lines.append("")
    lines.append(f"## ⬜ 不需正视（{len(no)} 个）—— 按原有提示词正常出图；想改随时说")
    for r in no:
        lines.append(f"- {r['name']} —— 主体：{r['subject']}；{r['reason']}")
    if bad:
        lines.append("")
        lines.append(f"## ⚠️ 评判失败/存疑（{len(bad)} 个）")
        for r in bad:
            lines.append(f"- {r['name']} —— {r['reason']}")
    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"\n审计清单已写入：{out_path}")
    print(f"  建议正视 {len(yes)} / 不需正视 {len(no)}" + (f" / 失败 {len(bad)}" if bad else ""))
    return 0


def rewrite(docx_path: Path, targets: list, out_path: Path):
    scenes = read_scene_docx(docx_path)
    missing = [t for t in targets if t not in scenes]
    if missing:
        print(f"!! 这些场景在 docx 里找不到：{missing}")
        print(f"   现有场景：{list(scenes.keys())}")
        return 2

    # 读原 docx 的全部 sections，替换目标场景的 content，其余原样保留
    doc = Document(str(docx_path))
    sections, cur_name = [], None
    buf_heading, buf_text = None, []
    for p in doc.paragraphs:
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            if buf_heading is not None:
                sections.append({"heading": buf_heading, "content": "\n".join(buf_text).strip()})
            buf_heading, buf_text = p.text.strip(), []
            m = HEAD_RE.match(p.text.strip())
            cur_name = m.group("name").strip() if m else None
        elif buf_heading is not None and p.text.strip():
            buf_text.append(p.text.strip())
    if buf_heading is not None:
        sections.append({"heading": buf_heading, "content": "\n".join(buf_text).strip()})

    targets = list(dict.fromkeys(targets))
    print(f"待改写场景 {len(targets)} 个：{targets}\n")
    changed, changed_scenes = 0, set()
    for t in targets:
        rec = scenes[t]
        keys = [k for k, v in rec["variants"].items() if v]
        # 逐版改写：每版一次调用，份数天然对齐，单版失败只保留该版原文
        for k in keys:
            orig = rec["variants"][k]
            try:
                nv = call_ds_api(REWRITE_SYSTEM, f"场景名：{t}\n\n{orig}",
                                 temperature=0.3, max_tokens=1800).strip()
            except Exception as e:
                print(f"  [警告] {t} 方案{k} —— 改写调用失败，保留原文: {e!r}")
                continue
            nv = re.sub(r"^方\s*案\s*[一二三四1234]\s*[:：]\s*", "", nv).strip()   # 洗掉模型自加的方案号
            if len(nv) < 60:
                print(f"  [警告] {t} 方案{k} —— 改写结果过短（{len(nv)} 字），保留原文")
                continue
            if not ("正对" in nv or "一点透视" in nv):
                print(f"  [注意] {t} 方案{k} —— 改写结果里没找到正对特征词，请人工复核")
            for s in sections:
                if s["heading"].startswith(f"{t} - 方案{k}（"):
                    s["content"] = nv
                    break
            changed += 1
            changed_scenes.add(t)
            print(f"  {t} 方案{k} ✓")

    if changed:
        title = next((p.text.strip() for p in doc.paragraphs
                      if (p.style.name or "").lower().startswith("title") and p.text.strip()),
                     "场景提示词")
        save_text_docx(title, sections, str(out_path))
        print(f"\n正视版已保存：{out_path}（改写 {changed} 版 / {len(changed_scenes)} 个场景，其余原样复制）")
    else:
        print("\n没有版本被改写，不写文件。")
    return 0


def main():
    ap = argparse.ArgumentParser(description="场景资产图正视评判与改写")
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("audit", help="逐场景评判是否需要正视，落盘审计清单")
    a.add_argument("--docx", required=True)
    a.add_argument("--out", default="")
    r = sub.add_parser("rewrite", help="把指定场景的三版改写成正视版（另存新 docx）")
    r.add_argument("--docx", required=True)
    r.add_argument("--scenes", required=True, help="逗号分隔的场景名")
    r.add_argument("--out", default="")
    args = ap.parse_args()

    docx_path = Path(args.docx)
    if not docx_path.exists():
        print(f"!! docx 不存在：{docx_path}")
        return 2
    if args.cmd == "audit":
        out = Path(args.out) if args.out else docx_path.with_name(docx_path.stem + "_正视审计.md")
        return audit(docx_path, out)
    targets = [s.strip() for s in args.scenes.split(",") if s.strip()]
    out = Path(args.out) if args.out else docx_path.with_name(docx_path.stem + "_正视版.docx")
    return rewrite(docx_path, targets, out)


if __name__ == "__main__":
    time.sleep(0)
    sys.exit(main())
