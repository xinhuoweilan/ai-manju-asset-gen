from pathlib import Path

from docx import Document


def read_docx_text(file_path: str) -> str:
    doc = Document(file_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # 剧本偶尔会放在表格里，一并取上（当前剧本无表格，不影响结果）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts)


def read_txt_text(file_path: str) -> str:
    # utf-8-sig 能同时吃「有 BOM」和「无 BOM」的 UTF-8；再兜一层 GBK
    for encoding in ["utf-8-sig", "gbk"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise RuntimeError(f"无法解码文件：{file_path}")


def read_script(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".docx":
        return read_docx_text(file_path)
    elif suffix == ".txt":
        return read_txt_text(file_path)
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")


def list_candidates(input_dir: str) -> list:
    """列出 input 目录里所有可用剧本，.txt 优先、.docx 次之，各自按名称排序。"""
    p = Path(input_dir)
    out = []
    for pattern in ("*.txt", "*.docx"):
        for f in sorted(p.glob(pattern)):
            if f.name.startswith("~$") or f.name.startswith("."):
                continue
            out.append(str(f))
    return out


def find_input_file(input_dir: str):
    """默认取第一个候选（.txt 优先，其次 .docx）。

    注意这个优先级本身是个坑：同一目录里若同时放了「节选.txt」和「全本.docx」，
    默认会取 .txt，导致 --scene 1-20 静默只处理 .txt 覆盖的那几集。
    main.py 会在候选多于一个时打印清单提醒；要指定文件请用 --input。
    """
    candidates = list_candidates(input_dir)
    return candidates[0] if candidates else None
