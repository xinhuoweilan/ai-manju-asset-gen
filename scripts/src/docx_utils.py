from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path


def save_text_docx(title: str, sections: list, output_path: str):
    if not sections:
        # 空列表会把上一批产物冲成一个「只有标题」的空文档 —— 直接不写、保留原文件。
        print(f"[跳过] 没有可写入的内容，保留原文件不动：{output_path}")
        return None

    doc = Document()

    title_paragraph = doc.add_heading(title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in sections:
        doc.add_heading(section["heading"], level=1)

        content = section["content"]
        if isinstance(content, list):
            content = "\n\n".join(content)

        paragraph = doc.add_paragraph(content)
        for run in paragraph.runs:
            run.font.size = Pt(11)

        doc.add_paragraph()

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def save_scene_table_docx(title: str, rows: list, output_path: str):
    doc = Document()

    title_paragraph = doc.add_heading(title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("场景—集数对照表（制表符分隔便于复制）")

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "场景"
    hdr_cells[1].text = "集数"

    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row[0]
        cells[1].text = "\t".join(row[1:]) if len(row) > 1 else ""

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path

